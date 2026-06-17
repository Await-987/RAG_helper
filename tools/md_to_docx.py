"""
将 Markdown 主稿转换为 docx 文件。
使用 python-docx 直接渲染，无需 pandoc。
支持: 标题层级、段落、有序/无序列表、表格、引用块、代码块、行内格式。
"""
import re
import sys
from pathlib import Path
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


def set_cell_border(cell):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_borders = OxmlElement('w:tcBorders')
    for edge in ('top', 'left', 'bottom', 'right'):
        b = OxmlElement(f'w:{edge}')
        b.set(qn('w:val'), 'single')
        b.set(qn('w:sz'), '4')
        b.set(qn('w:color'), '808080')
        tc_borders.append(b)
    tc_pr.append(tc_borders)


def set_run_chinese_font(run, font_name='宋体', size=10.5):
    run.font.name = font_name
    run.font.size = Pt(size)
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts')
        rPr.append(rFonts)
    rFonts.set(qn('w:eastAsia'), font_name)
    rFonts.set(qn('w:ascii'), 'Times New Roman')
    rFonts.set(qn('w:hAnsi'), 'Times New Roman')


def add_inline_runs(paragraph, text, base_size=10.5, base_font='宋体'):
    """处理行内的 **粗体**、*斜体*、`代码`、$公式$"""
    pattern = re.compile(r'(\*\*[^\*]+\*\*|\*[^\*]+\*|`[^`]+`|\$[^\$]+\$)')
    pos = 0
    for m in pattern.finditer(text):
        if m.start() > pos:
            r = paragraph.add_run(text[pos:m.start()])
            set_run_chinese_font(r, base_font, base_size)
        token = m.group()
        if token.startswith('**'):
            r = paragraph.add_run(token[2:-2])
            r.bold = True
            set_run_chinese_font(r, base_font, base_size)
        elif token.startswith('*'):
            r = paragraph.add_run(token[1:-1])
            r.italic = True
            set_run_chinese_font(r, base_font, base_size)
        elif token.startswith('`'):
            r = paragraph.add_run(token[1:-1])
            r.font.name = 'Consolas'
            r.font.size = Pt(base_size - 0.5)
            r.font.color.rgb = RGBColor(0xC7, 0x25, 0x4E)
        elif token.startswith('$'):
            r = paragraph.add_run(token[1:-1])
            r.italic = True
            r.font.name = 'Cambria Math'
            r.font.size = Pt(base_size)
        pos = m.end()
    if pos < len(text):
        r = paragraph.add_run(text[pos:])
        set_run_chinese_font(r, base_font, base_size)


def parse_table(lines, start_idx):
    """解析 Markdown 表格，返回 (rows, end_idx)"""
    rows = []
    i = start_idx
    while i < len(lines):
        ln = lines[i].rstrip()
        if not ln.startswith('|'):
            break
        if re.match(r'^\|[\s\-:|]+\|$', ln):
            i += 1
            continue
        cells = [c.strip() for c in ln.strip('|').split('|')]
        rows.append(cells)
        i += 1
    return rows, i


def render_md_to_docx(md_path: Path, docx_path: Path):
    text = md_path.read_text(encoding='utf-8')
    lines = text.split('\n')

    doc = Document()
    style = doc.styles['Normal']
    style.font.name = 'Times New Roman'
    style.font.size = Pt(10.5)
    rPr = style.element.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts')
        rPr.append(rFonts)
    rFonts.set(qn('w:eastAsia'), '宋体')

    section = doc.sections[0]
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)

    i = 0
    in_frontmatter = False
    in_code_block = False
    code_buffer = []

    while i < len(lines):
        line = lines[i]

        if i == 0 and line.strip() == '---':
            in_frontmatter = True
            i += 1
            continue
        if in_frontmatter:
            if line.strip() == '---':
                in_frontmatter = False
            i += 1
            continue

        if line.startswith('```'):
            if in_code_block:
                in_code_block = False
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Cm(0.5)
                p.paragraph_format.space_before = Pt(3)
                p.paragraph_format.space_after = Pt(3)
                r = p.add_run('\n'.join(code_buffer))
                r.font.name = 'Consolas'
                r.font.size = Pt(9)
                code_buffer = []
            else:
                in_code_block = True
            i += 1
            continue
        if in_code_block:
            code_buffer.append(line)
            i += 1
            continue

        stripped = line.strip()

        if not stripped:
            i += 1
            continue

        if stripped.startswith('# '):
            doc.add_page_break() if i > 0 else None
            h = doc.add_heading(stripped[2:].strip(), level=1)
            for r in h.runs:
                set_run_chinese_font(r, '黑体', 18)
            i += 1
            continue
        if stripped.startswith('## '):
            h = doc.add_heading(stripped[3:].strip(), level=2)
            for r in h.runs:
                set_run_chinese_font(r, '黑体', 15)
            i += 1
            continue
        if stripped.startswith('### '):
            h = doc.add_heading(stripped[4:].strip(), level=3)
            for r in h.runs:
                set_run_chinese_font(r, '黑体', 13)
            i += 1
            continue
        if stripped.startswith('#### '):
            h = doc.add_heading(stripped[5:].strip(), level=4)
            for r in h.runs:
                set_run_chinese_font(r, '黑体', 11.5)
            i += 1
            continue

        if stripped == '---':
            doc.add_page_break()
            i += 1
            continue

        if stripped.startswith('|') and i + 1 < len(lines) and re.match(r'^\|[\s\-:|]+\|$', lines[i + 1].strip()):
            rows, new_i = parse_table(lines, i)
            if rows:
                tbl = doc.add_table(rows=len(rows), cols=len(rows[0]))
                tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
                tbl.autofit = True
                for r_idx, row in enumerate(rows):
                    for c_idx, cell_text in enumerate(row):
                        if c_idx >= len(rows[0]):
                            break
                        cell = tbl.rows[r_idx].cells[c_idx]
                        cell.text = ''
                        p = cell.paragraphs[0]
                        add_inline_runs(p, cell_text, base_size=9.5)
                        if r_idx == 0:
                            for r in p.runs:
                                r.bold = True
                        set_cell_border(cell)
            i = new_i
            continue

        if stripped.startswith('> '):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.8)
            p.paragraph_format.space_after = Pt(6)
            r = p.add_run(stripped[2:].strip())
            r.italic = True
            r.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
            set_run_chinese_font(r, '楷体', 10.5)
            i += 1
            continue

        m_ol = re.match(r'^(\d+)\.\s+(.*)$', stripped)
        m_ul = re.match(r'^[-*]\s+(.*)$', stripped)
        if m_ol:
            p = doc.add_paragraph(style='List Number')
            add_inline_runs(p, m_ol.group(2))
            i += 1
            continue
        if m_ul:
            p = doc.add_paragraph(style='List Bullet')
            add_inline_runs(p, m_ul.group(1))
            i += 1
            continue

        p = doc.add_paragraph()
        p.paragraph_format.first_line_indent = Cm(0.74)
        p.paragraph_format.line_spacing = 1.5
        add_inline_runs(p, stripped)
        i += 1

    doc.save(str(docx_path))
    print(f"[OK] saved: {docx_path}  ({docx_path.stat().st_size/1024:.1f} KB)")


if __name__ == '__main__':
    md = Path(sys.argv[1] if len(sys.argv) > 1 else '/home/ubuntu/rag_project/rag/技术报告_完整版.md')
    out = Path(sys.argv[2] if len(sys.argv) > 2 else '/home/ubuntu/rag_project/rag/技术报告_完整版.docx')
    render_md_to_docx(md, out)
